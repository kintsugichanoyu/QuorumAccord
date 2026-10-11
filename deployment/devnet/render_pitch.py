"""Create an editable, linked PowerPoint from the evidence-backed pitch data."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
OUT=Path(__file__).parent/'evidence'/'v2'
data=json.loads((OUT/'pitch-slides.json').read_text())
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
BG,INK,MUTED,GREEN='F3F5EE','1E3022','728069','9ABB68'
def text(slide,x,y,w,h,value,size=18,color=INK,font='Aptos',bold=False):
 shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=shape.text_frame;frame.word_wrap=True;frame.margin_left=0;frame.margin_right=0
 for i,line in enumerate(value.split('\n')):
  p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.text=line;p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(color);p.font.name=font;p.font.bold=bold;p.space_after=Pt(9)
 return shape
for i,s in enumerate(data['slides']):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 text(slide,.85,.5,14,.4,s['kicker'],11,GREEN,bold=True)
 text(slide,.85,1.12,14.3,1.25,s['title'],54 if s['kind']=='cover' else 42,bold=True)
 text(slide,.85,2.55,14.2,.9,s['subtitle'],18,MUTED)
 points=s['points'];kind=s['kind'];top=3.6
 if kind in ('three','metrics','cover','allocation','pipeline'):
  cols=len(points);gap=.25;width=(14.3-gap*(cols-1))/cols
  for j,p in enumerate(points):
   x=.85+j*(width+gap);shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(top),Inches(width),Inches(2.75));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('D8ECB0' if kind=='cover' else 'E9EFDF');shape.line.fill.background()
   text(slide,x+.2,top+.22,width-.4,.4,f'{j+1:02}',11,GREEN)
   text(slide,x+.2,top+.65,width-.4,1.9,p,27 if kind=='allocation' else 23 if kind=='metrics' else 18)
 else:
  row=min(.72,3.9/max(1,len(points))) if kind=='list' else 1.1
  for j,p in enumerate(points):
   y=top+j*row;text(slide,.85,y,.35,row,f'{j+1:02}',10,GREEN)
   text(slide,1.35,y,13.8,row,p,14 if len(''.join(points))>850 else 17,font='Aptos Mono' if kind=='snapshot' else 'Aptos')
 for j,l in enumerate(s.get('links',[])):
  shape=text(slide,.85+j*4.7,7.5,4.4,.5,l['label']+' ↗',12,INK)
  for p in shape.text_frame.paragraphs:
   for run in p.runs:run.hyperlink.address=l['url'] if not l['url'].startswith('/') else 'http://localhost:18880'+l['url']
 text(slide,.85,8.45,13,.3,'QUORUM · USDC / DEVNET · CAPTURED '+data['built_at'][:10],9,MUTED)
 text(slide,14,8.45,1,.3,f'{i+1:02} / {len(data["slides"])}',9,MUTED)
 slide.notes_slide.notes_text_frame.text='\n'.join(s['notes'])
prs.save(OUT/'pitch.pptx')
(Path(__file__).parents[2]/'demo-site'/'dist'/'pitch.pptx').write_bytes((OUT/'pitch.pptx').read_bytes())
print('Wrote',len(prs.slides),'editable PowerPoint slides.')
